"""build_docx.py — render a journal manuscript to .docx via python-docx.

Protocol: JSON on stdin, output path as the only argument.
    {"title": "...", "paperKind": "...", "venue": "...",
     "sections": [{"heading": "...", "body": "..."}],
     "references": [{"title": "...", "authors": "[...]", "year": 2024, ...}]}

Prints one JSON line {"ok": true, "path": "..."} on success.
"""
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def set_cjk_font(run, name="宋体", size=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def add_paragraph(doc, text, size=10.5, first_indent=True):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    run = paragraph.add_run(text)
    set_cjk_font(run, size=size)


def main():
    out_path = sys.argv[1]
    # Always decode as UTF-8: the Rust caller writes UTF-8 bytes, and Windows
    # consoles default stdin to GBK, which would corrupt CJK into control
    # characters that lxml rejects.
    raw = sys.stdin.buffer.read().decode("utf-8")
    payload = json.loads(raw)
    doc = Document()

    title = str(payload.get("title") or "未命名论文")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    set_cjk_font(run, size=18)

    meta = " · ".join(
        str(item) for item in (payload.get("paperKind"), payload.get("venue")) if item
    )
    if meta:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(meta)
        set_cjk_font(run, size=10.5)

    sections = payload.get("sections") or []
    for section in sections:
        heading = str(section.get("heading") or "").strip()
        body = str(section.get("body") or "").strip()
        if not heading and not body:
            continue
        if heading:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(heading)
            run.bold = True
            set_cjk_font(run, size=14)
        for block in body.split("\n\n"):
            text = block.strip()
            if text:
                add_paragraph(doc, text)

    references = payload.get("references") or []
    if references:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("参考文献")
        run.bold = True
        set_cjk_font(run, size=14)
        for index, ref in enumerate(references, start=1):
            authors = ref.get("authors") or "[]"
            try:
                author_list = json.loads(authors) if isinstance(authors, str) else authors
                author_text = ", ".join(str(name) for name in author_list)
            except Exception:
                author_text = str(authors)
            year = ref.get("year")
            venue = ref.get("venue") or ""
            title_text = ref.get("title") or ""
            doi = ref.get("doi") or ""
            line = f"[{index}] {author_text}. {title_text}[J]. {venue}"
            if year:
                line += f", {year}"
            if doi:
                line += f". DOI: {doi}"
            line += "."
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(line)
            set_cjk_font(run, size=9)

    doc.save(out_path)
    print(json.dumps({"ok": True, "path": out_path}, ensure_ascii=False))


if __name__ == "__main__":
    main()
