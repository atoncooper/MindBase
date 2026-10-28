"""数据集画像 sidecar（由桌面端 dataset.rs 调用）。

用法：
    python dataset_profile.py <文件路径>

支持 .csv（标准库 csv，自动识别 UTF-8 / GB18030 与分隔符）与
.xlsx（openpyxl，只读流式读取）。

输出协议（与 doc_extract.py / pptx_build.py 一致）：stdout 的**最后一行**
是唯一的判定 JSON，第三方库的告警可能先打印在上面。Rust 侧永远取最后
一个非空行解析，因此这里只 print 一行、不打印其它内容。

成功：
    {"ok": true, "rows": 120, "cols": 5,
     "columns": [{"name": "age", "type": "integer", "missing": 0,
                  "unique": 42, "min": 18, "max": 76, "mean": 34.12}, ...],
     "preview": [["张三", "34", ...], ...]}
失败：
    {"ok": false, "error": "中文原因"}

字段语义：
- type ∈ integer / number / date / string（整列可转数字即数值型；
  非空值全部形如日期即 date，否则 string）；
- missing = 空值单元格数（""、NA、N/A、null、None、NaN 都算缺失）；
- unique = 非空去重值个数；
- min / max / mean 只对数值列有值（date 列的 min/max 是 ISO 日期字符串，
  mean 为 null；非数值列 min/max 为字典序极值，仅作参考）；
- preview 是前 5 个数据行，单元格统一转成字符串（截断 160 字符）。

任何异常都降级为 {"ok": false, "error": ...}（进程退出码恒为 0），由
Rust 转成中文错误信息抛给前端。
"""

import csv
import io
import json
import math
import re
import statistics
import sys
from datetime import date, datetime
from pathlib import Path

# 视为「缺失」的文本标记（比较前统一转小写并 strip）。
MISSING_TOKENS = {"", "na", "n/a", "null", "none", "nan"}
# 画像预览保留的数据行数。
PREVIEW_ROWS = 5
# 形如 2024-01-02 / 2024/1/2 03:04:05 的日期文本。
DATE_RE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$")
# 带千分位的数字（Excel 导出的 CSV 常见），如 1,234 / -1,234.5。
THOUSANDS_RE = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")


def _is_missing(value):
    """单元格是否算缺失值。"""
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str):
        return value.strip().lower() in MISSING_TOKENS
    return False


def _cell_text(value, limit=400):
    """把任意单元格值转成一行短文本（用于 preview / 去重 / 字符串统计）。"""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        text = value.isoformat()
    elif isinstance(value, float):
        text = "%g" % value
    else:
        text = str(value)
    text = text.replace("\r", " ").replace("\n", " ").strip()
    return text[:limit]


def _to_number(value):
    """尝试转成 float；不是数字时返回 None。"""
    if isinstance(value, bool):
        # Excel 的 TRUE/FALSE 不当数值统计。
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _cell_text(value)
        if THOUSANDS_RE.match(text):
            text = text.replace(",", "")
        try:
            number = float(text)
        except ValueError:
            return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def _to_date(value):
    """尝试转成 date；不是日期时返回 None。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    if not DATE_RE.match(text):
        return None
    normalized = text.replace("/", "-").replace("T", " ")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(normalized, fmt).date()
        except ValueError:
            continue
    return None


def _describe(present):
    """推断一列的 type / min / max / mean（present 已剔除缺失值）。"""
    if not present:
        return "string", None, None, None

    numbers = [_to_number(value) for value in present]
    if all(number is not None for number in numbers):
        is_integer = all(float(number).is_integer() for number in numbers)
        minimum = min(numbers)
        maximum = max(numbers)
        mean = round(statistics.fmean(numbers), 6)
        if is_integer:
            return "integer", int(minimum), int(maximum), mean
        return "number", minimum, maximum, mean

    dates = [_to_date(value) for value in present]
    if all(moment is not None for moment in dates):
        iso = [moment.isoformat() for moment in dates]
        return "date", min(iso), max(iso), None

    texts = [_cell_text(value) for value in present]
    return "string", min(texts), max(texts), None


def _column_profile(name, cells):
    """单列画像：cells 含该列全部数据行（含缺失值）。"""
    present = [cell for cell in cells if not _is_missing(cell)]
    kind, minimum, maximum, mean = _describe(present)
    return {
        "name": name,
        "type": kind,
        "missing": len(cells) - len(present),
        "unique": len({_cell_text(cell) for cell in present}),
        "min": minimum,
        "max": maximum,
        "mean": mean,
    }


def _header_names(header_row):
    """表头转列名：空表头补 column_N，重名列加序号。"""
    names = []
    seen = {}
    for index, cell in enumerate(header_row):
        name = _cell_text(cell, 80) or "column_%d" % (index + 1)
        count = seen.get(name, 0) + 1
        seen[name] = count
        if count > 1:
            name = "%s_%d" % (name, count)
        names.append(name)
    return names


def _read_csv_rows(path):
    """读 CSV 全部行（先 UTF-8，失败回退 GB18030；自动识别分隔符）。"""
    text = None
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                text = handle.read()
            break
        except UnicodeDecodeError as err:
            last_error = err
    if text is None:
        raise ValueError("无法识别 CSV 编码（已尝试 UTF-8 / GB18030）：%s" % last_error)
    if not text.strip():
        return []
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        # 单列 CSV 或样本不足时用逗号兜底。
        delimiter = ","
    # 用 StringIO 而不是 splitlines()：带引号的多行单元格不会被拆坏。
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_xlsx_rows(path):
    """用 openpyxl 只读模式读取活动工作表（含表头的第一行起）。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("缺少 openpyxl 依赖，无法读取 xlsx（请重新导入以安装依赖）")

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = list(row)
            while cells and cells[-1] is None:
                cells.pop()
            rows.append(cells)
        return rows
    finally:
        workbook.close()


def _profile(header, data_rows):
    """整表画像：逐列统计 + 前 N 行预览。"""
    cols = len(header)
    columns = [[] for _ in range(cols)]
    preview = []
    for row in data_rows:
        cells = list(row[:cols])
        if len(cells) < cols:
            cells.extend([None] * (cols - len(cells)))
        for index in range(cols):
            columns[index].append(cells[index])
        if len(preview) < PREVIEW_ROWS:
            preview.append([_cell_text(value, 160) for value in cells])
    return {
        "rows": len(data_rows),
        "cols": cols,
        "columns": [
            _column_profile(name, cells) for name, cells in zip(header, columns)
        ],
        "preview": preview,
    }


def _profile_file(raw_path):
    """读文件并返回画像结果；格式不支持/为空时返回 ok:false。"""
    path = Path(raw_path)
    if not path.is_file():
        return {"ok": False, "error": "文件不存在：%s" % path}
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv_rows(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _read_xlsx_rows(path)
    elif suffix == ".xls":
        return {
            "ok": False,
            "error": "暂不支持旧版 .xls 格式，请在 Excel 中另存为 .xlsx 后再导入",
        }
    else:
        return {
            "ok": False,
            "error": "不支持的数据格式：%s（仅支持 csv / xlsx）" % (suffix or "未知"),
        }

    # 丢弃整行为空的行（Excel 尾部常有空行），避免把行数算多。
    rows = [row for row in rows if any(not _is_missing(cell) for cell in row)]
    if not rows:
        return {"ok": False, "error": "数据集为空（没有任何数据行）"}

    header_row = list(rows[0])
    while header_row and _is_missing(header_row[-1]):
        header_row.pop()
    header = _header_names(header_row)
    if not header:
        return {"ok": False, "error": "数据集缺少表头（第一行为空）"}

    profiled = _profile(header, rows[1:])
    result = {"ok": True}
    result.update(profiled)
    return result


def main(argv):
    if len(argv) < 2:
        return {"ok": False, "error": "缺少输入文件路径"}
    return _profile_file(argv[1])


if __name__ == "__main__":
    try:
        verdict = main(sys.argv)
    except Exception as err:  # noqa: BLE001 — sidecar 协议：任何异常降级为 ok:false
        verdict = {"ok": False, "error": "画像失败：%s" % err}
    print(json.dumps(verdict, ensure_ascii=False))
